"""Offline parser boundary tests; run with `python -m unittest test_parser`."""

import base64
import io
import json
import os
import socket
import tempfile
import threading
import time
import unittest
import zipfile
from unittest import mock

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pptx import Presentation

import parser


def as_bytes(save):
    stream = io.BytesIO()
    save(stream)
    return stream.getvalue()


class ParserTests(unittest.TestCase):
    def test_extracts_plain_text_mail_body(self):
        self.assertEqual(parser.extract("message.txt", b"\xef\xbb\xbfApproved travel\nPolicy"), "Approved travel\nPolicy")

    def test_extracts_docx_text(self):
        def save(out):
            doc = Document()
            doc.add_paragraph("Approved policy text")
            doc.save(out)

        self.assertIn("Approved policy text", parser.extract("policy.docx", as_bytes(save)))

    def test_rejects_docx_with_too_many_table_cells(self):
        def save(out):
            doc = Document()
            doc.add_table(rows=1, cols=2)
            doc.save(out)

        with mock.patch.object(parser, "MAX_CELLS", 1), self.assertRaisesRegex(
            parser.ExtractionError, "cell_limit"
        ):
            parser.extract("wide-table.docx", as_bytes(save))

    def test_rejects_docx_with_too_many_table_paragraphs(self):
        def save(out):
            doc = Document()
            doc.add_table(rows=1, cols=1).cell(0, 0).add_paragraph("another paragraph")
            doc.save(out)

        with mock.patch.object(parser, "MAX_PARAGRAPHS", 1), self.assertRaisesRegex(
            parser.ExtractionError, "paragraph_limit"
        ):
            parser.extract("many-paragraphs.docx", as_bytes(save))

    def test_extracts_pptx_text(self):
        def save(out):
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = "Quarterly plan"
            deck.save(out)

        self.assertIn("Quarterly plan", parser.extract("plan.pptx", as_bytes(save)))

    def test_extracts_xlsx_values(self):
        def save(out):
            book = Workbook()
            book.active["A1"] = "Budget"
            book.active["B1"] = 120
            book.save(out)

        text = parser.extract("budget.xlsx", as_bytes(save))
        self.assertIn("Budget", text)
        self.assertIn("120", text)

    def test_rejects_unsupported_type(self):
        with self.assertRaisesRegex(parser.ExtractionError, "unsupported_file_type"):
            parser.extract("archive.zip", b"PK")

    def test_rejects_zip_expansion_ratio(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("payload.xml", "a" * 100_000)
        with self.assertRaisesRegex(parser.ExtractionError, "archive_expansion_limit"):
            parser.check_zip(stream.getvalue())

    def test_rejects_corrupt_documents(self):
        for filename, data in (("broken.pdf", b"%PDF-not-a-pdf"), ("broken.docx", b"not a zip")):
            with self.subTest(filename=filename), self.assertRaisesRegex(parser.ExtractionError, "invalid_document"):
                parser.extract(filename, data)

    def test_maps_zip_library_failures_to_invalid_document(self):
        with mock.patch.object(parser.zipfile, "ZipFile", side_effect=OSError("malformed central directory")):
            with self.assertRaisesRegex(parser.ExtractionError, "invalid_document"):
                parser.extract("broken.docx", b"PK")

    def test_rejects_encrypted_pdf(self):
        stream = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.encrypt("user", "owner")
        writer.write(stream)
        with self.assertRaisesRegex(parser.ExtractionError, "encrypted_documents_unsupported"):
            parser.extract("encrypted.pdf", stream.getvalue())

    def test_rejects_zip_expansion_and_entry_count_limits(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_STORED) as archive:
            archive.writestr("one.xml", "x" * 33)
        with mock.patch.object(parser, "MAX_ZIP_EXPANDED", 32), self.assertRaisesRegex(
            parser.ExtractionError, "archive_expansion_limit"
        ):
            parser.check_zip(stream.getvalue())

        with mock.patch.object(parser, "MAX_ZIP_ENTRIES", 0), self.assertRaisesRegex(
            parser.ExtractionError, "archive_entry_limit"
        ):
            parser.check_zip(stream.getvalue())

    def test_rejects_text_and_pdf_page_limits(self):
        with mock.patch.object(parser, "MAX_TEXT", 4), self.assertRaisesRegex(
            parser.ExtractionError, "document_text_limit"
        ):
            parser.extract("large.txt", b"12345")

        reader = mock.Mock(is_encrypted=False, pages=[None, None])
        with mock.patch.object(parser, "MAX_PDF_PAGES", 1), mock.patch.object(parser, "PdfReader", return_value=reader):
            with self.assertRaisesRegex(parser.ExtractionError, "page_limit"):
                parser.extract("many-pages.pdf", b"%PDF-fixture")

    def test_rejects_input_over_limit(self):
        with self.assertRaisesRegex(parser.ExtractionError, "input_size_limit"):
            parser.extract("file.pdf", b"%PDF-" + b"x" * parser.MAX_INPUT)

    def test_extraction_deadline_stops_slow_document(self):
        with mock.patch.object(parser, "PdfReader", side_effect=lambda *_args, **_kwargs: time.sleep(1)):
            with self.assertRaisesRegex(parser.ExtractionError, "processing_time_limit"):
                parser.extract_with_deadline("document.pdf", b"%PDF-fixture", timeout=0.02)

    def test_socket_handler_returns_503_on_extraction_timeout(self):
        with tempfile.TemporaryDirectory() as directory:
            socket_path = os.path.join(directory, "parser.sock")
            with parser.UnixHTTPServer(socket_path, parser.Handler) as server:
                payload = json.dumps({
                    "filename": "slow.pdf",
                    "contentBase64": base64.b64encode(b"%PDF-slow").decode("ascii"),
                }).encode("utf-8")
                request = (
                    b"POST /v1/parse HTTP/1.1\r\nHost: parser\r\nContent-Type: application/json\r\n"
                    + f"Content-Length: {len(payload)}\r\nConnection: close\r\n\r\n".encode("ascii")
                    + payload
                )
                response = []

                def send_request():
                    with socket.socket(socket.AF_UNIX, socket.SOCK_STREAM) as client:
                        client.connect(socket_path)
                        client.sendall(request)
                        chunks = []
                        while chunk := client.recv(4096):
                            chunks.append(chunk)
                        response.append(b"".join(chunks))

                client_thread = threading.Thread(target=send_request)
                client_thread.start()
                with mock.patch.object(parser, "MAX_EXTRACTION_SECONDS", 0.02), mock.patch.object(
                    parser, "PdfReader", side_effect=lambda *_args, **_kwargs: time.sleep(1)
                ):
                    server.handle_request()
                client_thread.join(timeout=2)
                self.assertFalse(client_thread.is_alive(), "parser socket client did not finish")
                headers, separator, body = response[0].partition(b"\r\n\r\n")
                self.assertTrue(separator)
                self.assertIn(b" 503 ", headers.split(b"\r\n", 1)[0])
                self.assertEqual(json.loads(body), {"error": "processing_time_limit"})


if __name__ == "__main__":
    unittest.main()
