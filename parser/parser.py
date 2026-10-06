"""Bounded, offline extraction for supported Office documents and PDFs."""

from __future__ import annotations

import base64
import binascii
import io
import json
import os
import re
import signal
import socketserver
import zipfile
from http.server import BaseHTTPRequestHandler
from typing import Iterable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

SOCKET = os.environ.get("PARSER_SOCKET", "/run/iqkb/parser.sock")
MAX_INPUT = 24 * 1024 * 1024
MAX_TEXT = 2_000_000
MAX_ZIP_EXPANDED = 128 * 1024 * 1024
MAX_ZIP_RATIO = 100
MAX_ZIP_ENTRIES = 20_000
MAX_PDF_PAGES = 200
MAX_SLIDES = 300
MAX_SHEETS = 100
MAX_CELLS = 100_000
MAX_PARAGRAPHS = 50_000
MAX_EXTRACTION_SECONDS = 20.0
SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".txt"}


class ExtractionError(ValueError):
    pass


# Escape parser-library `except Exception` blocks so the watchdog stays distinct.
class ExtractionDeadline(BaseException):
    pass


def extraction_deadline(_signum: int, _frame: object) -> None:
    raise ExtractionDeadline()


def bounded_text(parts: Iterable[str]) -> str:
    chunks: list[str] = []
    size = 0
    for part in parts:
        clean = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(part))
        if not clean.strip():
            continue
        if size + len(clean) + 1 > MAX_TEXT:
            raise ExtractionError("document_text_limit")
        chunks.append(clean)
        size += len(clean) + 1
    return "\n".join(chunks)


def check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            expanded = sum(item.file_size for item in infos)
            compressed = max(1, sum(item.compress_size for item in infos))
            if expanded > MAX_ZIP_EXPANDED or expanded / compressed > MAX_ZIP_RATIO:
                raise ExtractionError("archive_expansion_limit")
            if len(infos) > MAX_ZIP_ENTRIES:
                raise ExtractionError("archive_entry_limit")
            if any(item.flag_bits & 1 for item in infos):
                raise ExtractionError("encrypted_documents_unsupported")
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("invalid_document") from exc


def extract(filename: str, data: bytes) -> str:
    suffix = os.path.splitext(filename.lower())[1]
    if suffix not in SUPPORTED:
        raise ExtractionError("unsupported_file_type")
    if not data or len(data) > MAX_INPUT:
        raise ExtractionError("input_size_limit")

    if suffix == ".txt":
        return bounded_text([data.decode("utf-8-sig", errors="replace")])

    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ExtractionError("invalid_document")
        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ExtractionError("encrypted_documents_unsupported")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ExtractionError("page_limit")
            return bounded_text(page.extract_text() or "" for page in reader.pages)
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError("invalid_document") from exc

    check_zip(data)
    stream = io.BytesIO(data)
    try:
        if suffix == ".docx":
            document = Document(stream)
            parts: list[str] = []
            if len(document._element.xpath(".//w:tc")) > MAX_CELLS:
                raise ExtractionError("cell_limit")
            if len(document._element.xpath(".//w:p")) > MAX_PARAGRAPHS:
                raise ExtractionError("paragraph_limit")
            paragraphs = document.paragraphs
            parts.extend(p.text for p in paragraphs)
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return bounded_text(parts)

        if suffix == ".pptx":
            presentation = Presentation(stream)
            if len(presentation.slides) > MAX_SLIDES:
                raise ExtractionError("slide_limit")
            parts = []
            for slide in presentation.slides:
                parts.extend(shape.text for shape in slide.shapes if shape.has_text_frame)
                for shape in slide.shapes:
                    if shape.has_table:
                        parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            return bounded_text(parts)

        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            if len(workbook.worksheets) > MAX_SHEETS:
                raise ExtractionError("sheet_limit")
            parts = []
            cells = 0
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) > MAX_CELLS or (sheet.max_column or 0) > MAX_CELLS:
                    raise ExtractionError("cell_limit")
                parts.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    if cells > MAX_CELLS:
                        raise ExtractionError("cell_limit")
                    values = [str(value) for value in row if value is not None]
                    if values:
                        parts.append(" | ".join(values))
            return bounded_text(parts)
        finally:
            workbook.close()
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("invalid_document") from exc


def extract_with_deadline(filename: str, data: bytes, timeout: float | None = None) -> str:
    if timeout is None:
        timeout = MAX_EXTRACTION_SECONDS
    if timeout <= 0:
        raise ValueError("timeout must be positive")
    if signal.getitimer(signal.ITIMER_REAL) != (0.0, 0.0):
        raise RuntimeError("parser extraction deadline is already active")
    previous_handler = signal.signal(signal.SIGALRM, extraction_deadline)
    signal.setitimer(signal.ITIMER_REAL, timeout)
    try:
        return extract(filename, data)
    except ExtractionDeadline as exc:
        raise ExtractionError("processing_time_limit") from exc
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous_handler)


class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def do_POST(self) -> None:
        if self.path != "/v1/parse":
            self.respond(404, {"error": "not_found"})
            return
        if self.headers.get_content_type() != "application/json":
            self.respond(415, {"error": "content_type_required"})
            return
        try:
            length = int(self.headers.get("Content-Length", "-1"))
        except ValueError:
            length = -1
        if length < 0 or length > MAX_INPUT * 4 // 3 + 16_384:
            self.respond(413, {"error": "input_size_limit"})
            return
        try:
            payload = json.loads(self.rfile.read(length))
            if not isinstance(payload, dict) or set(payload) != {"filename", "contentBase64"}:
                raise ExtractionError("invalid_request")
            filename = payload["filename"]
            encoded = payload["contentBase64"]
            if not isinstance(filename, str) or len(filename) > 255 or not isinstance(encoded, str):
                raise ExtractionError("invalid_request")
            data = base64.b64decode(encoded, validate=True)
            text = extract_with_deadline(filename, data)
            self.respond(200, {"text": text})
        except (UnicodeDecodeError, json.JSONDecodeError, binascii.Error, ExtractionError) as exc:
            code = str(exc) if isinstance(exc, ExtractionError) else "invalid_request"
            if code == "processing_time_limit":
                status = 503
            elif code == "unsupported_file_type":
                status = 415
            elif code.endswith("limit"):
                status = 413
            else:
                status = 400
            self.respond(status, {"error": code or "invalid_request"})

    def respond(self, status: int, payload: dict[str, object]) -> None:
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Connection", "close")
        self.end_headers()
        self.wfile.write(body)
        self.close_connection = True

    def log_message(self, *_args: object) -> None:
        return


class UnixHTTPServer(socketserver.UnixStreamServer):
    allow_reuse_address = True
    request_queue_size = 8
    timeout = 30

    def server_bind(self) -> None:
        if os.path.exists(self.server_address):
            os.unlink(self.server_address)
        super().server_bind()
        os.chmod(self.server_address, 0o600)

    def get_request(self):
        sock, address = super().get_request()
        sock.settimeout(30)
        return sock, address


if __name__ == "__main__":
    os.makedirs(os.path.dirname(SOCKET), mode=0o700, exist_ok=True)
    with UnixHTTPServer(SOCKET, Handler) as server:
        server.serve_forever(poll_interval=0.5)
